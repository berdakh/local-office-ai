"""Build a simple letterhead .docx (header with name and address, footer with contact line).

Churches usually have a Word letterhead already; this is the fallback so every
install has a template on day one. Replace the file with theirs when available.
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def build_letterhead(
    path: Path,
    name: str,
    address: str = "",
    phone: str = "",
    email: str = "",
    website: str = "",
    accent_hex: str = "2F6F5E",
) -> Path:
    doc = docx.Document()
    section = doc.sections[0]

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(accent_hex)
    if address:
        p2 = header.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(address)
        r2.font.size = Pt(10)

    contact = "  |  ".join(x for x in (phone, email, website) if x)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(contact)
    fr.font.size = Pt(9)

    # Body stays empty; write_docx fills it.
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
