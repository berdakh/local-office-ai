import docx

from app.render import build_letterhead, write_docx


def test_letterhead_header_survives_rendering(tmp_path):
    tpl = build_letterhead(
        tmp_path / "letterhead.docx",
        "Grace Community Church",
        address="123 Main St",
        phone="(281) 555-0100",
        email="office@example.org",
    )
    out = write_docx("# Discussion Guide\n\nHello.\n- one\n1. first", tmp_path / "out.docx", tpl)
    d = docx.Document(str(out))
    header = "\n".join(p.text for p in d.sections[0].header.paragraphs)
    footer = "\n".join(p.text for p in d.sections[0].footer.paragraphs)
    assert "Grace Community Church" in header and "123 Main St" in header
    assert "(281) 555-0100" in footer and "office@example.org" in footer
    body = [p.text for p in d.paragraphs]
    assert body[0] == "Discussion Guide" and "Hello." in body
    styles = {p.style.name for p in d.paragraphs}
    assert "List Bullet" in styles and "List Number" in styles
